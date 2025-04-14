from dotenv import load_dotenv
import os
import oss2
import re
import openpyxl
import json
import matplotlib.pyplot as plt
import streamlit as st

# 加载 .env 文件
load_dotenv()

# 读取环境变量
access_key_id = os.getenv('ACCESS_KEY_ID')
access_key_secret = os.getenv('ACCESS_KEY_SECRET')

# 初始化 OSS 客户端
auth = oss2.Auth(access_key_id, access_key_secret)
bucket = oss2.Bucket(auth, 'https://oss-cn-shanghai.aliyuncs.com', 'checkbill')  # 修改为正确的 endpoint

# 获取 hongxin 目录下的文件列表
file_list = []
for obj in oss2.ObjectIterator(bucket, prefix='hongxin/'):
    if not obj.key.endswith('/'):  # 排除目录
        file_list.append(obj.key)

# 处理每个文件
result = []
for file_key in file_list:
    # 下载文件到本地
    local_file_path = os.path.join("temp", os.path.basename(file_key))
    os.makedirs("temp", exist_ok=True)  # 确保 temp 目录存在
    bucket.get_object_to_file(file_key, local_file_path)

    # 使用正则表达式提取年和月
    match = re.search(r'(\d{4})年(\d{1,2})月', file_key)
    if match:
        year = match.group(1)
        month = match.group(2)

        # 使用 openpyxl 打开文件
        wb = openpyxl.load_workbook(local_file_path)
        ws = wb.active

        # 读取 D10、D11、D12 单元格的数据
        d10 = ws['D10'].value
        d11 = ws['D11'].value
        d12 = ws['D12'].value

        # 处理带有 "=" 符号的数据
        def clean_value(value):
            if isinstance(value, str) and value.startswith('='):
                return float(value.lstrip('='))  # 去掉 "=" 并转换为数值
            return float(value) if value is not None else 0  # 确保返回数值

        d10 = clean_value(d10)
        d11 = clean_value(d11)
        d12 = clean_value(d12)

        # 保存结果
        result.append({
            "年": year,
            "月": month,
            "D10": d10,
            "D11": d11,
            "D12": d12
        })

    # 删除临时文件
    os.remove(local_file_path)

# 将结果保存为 JSON 文件
result_json_path = os.path.join("result.json")
with open(result_json_path, "w", encoding="utf-8") as f:
    json.dump(result, f, ensure_ascii=False, indent=4)

# 绘制折线图
# 按时间顺序排序
result_sorted = sorted(result, key=lambda x: (int(x['年']), int(x['月'])))  # 将年和月转换为整数排序
years_months = [f"{item['年']}-{item['月']}" for item in result_sorted]
d10_values = [item['D10'] for item in result_sorted]
d11_values = [item['D11'] for item in result_sorted]
d12_values = [item['D12'] for item in result_sorted]

# 设置中文字体
plt.rcParams['font.sans-serif'] = ['SimHei']  # 使用黑体
plt.rcParams['axes.unicode_minus'] = False  # 解决负号显示问题

plt.figure(figsize=(10, 6))
plt.plot(years_months, d10_values, label='物流部')
plt.plot(years_months, d11_values, label='事业部')
plt.plot(years_months, d12_values, label='宏信E采')
plt.xlabel('年月')
plt.ylabel('通话量')
plt.title('通话量随时间变化')
plt.legend()
plt.xticks(rotation=45)
plt.tight_layout()
chart_path = os.path.join("call_volume_chart.png")
plt.savefig(chart_path)

# 使用 Streamlit 显示结果
st.title("宏信外呼量统计")

# 显示折线图
st.subheader("通话量随时间变化")
st.image(chart_path)

# 显示详细数字
st.subheader("详细数字")
with st.expander("点击展开查看详细数字", expanded=False):  # 默认收起
    st.json(result)

st.write("处理完成，结果已保存到 result.json，图表已保存到 call_volume_chart.png")